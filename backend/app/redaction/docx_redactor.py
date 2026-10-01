from docx import Document

def redact_docx(input_path: str, output_path: str, redactions: list):
    try:
        doc = Document(input_path)
        for r in redactions:
            for para in doc.paragraphs:
                if r['value'] in para.text:
                    para.text = para.text.replace(r['value'], r['masked_value'])
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if r['value'] in cell.text:
                            cell.text = cell.text.replace(r['value'], r['masked_value'])
        doc.save(output_path)
    except Exception as e:
        print(f"DOCX Redact Error: {e}")
